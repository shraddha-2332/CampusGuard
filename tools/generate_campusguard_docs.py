from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from pathlib import Path


OUT_DIR = Path("deliverables")
PROJECT_NAME = "CampusGuard"
PROJECT_FULL_NAME = "CampusGuard AI Powered University Admission Engagement and Information Support System"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_padding(cell, value="120"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")


def style_table(table, widths=None):
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_padding(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[col_idx]
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)
            if row_idx == 0:
                set_cell_shading(cell, "1F4E79")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = None
            elif row_idx % 2 == 0:
                set_cell_shading(cell, "F4F8FB")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table, widths=widths)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_meta_table(doc, rows):
    add_table(
        doc,
        ["Field", "Value"],
        rows,
        widths=[Inches(1.8), Inches(4.7)],
    )


def setup_doc(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.size = Pt(size)
        styles[style_name].font.bold = True

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(title)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(10)
    r = sp.add_run(subtitle)
    r.font.name = "Aptos"
    r.font.size = Pt(12)
    r.bold = True
    return doc


def save(doc, filename):
    OUT_DIR.mkdir(exist_ok=True)
    path = OUT_DIR / filename
    doc.save(path)
    return path


def common_review_points():
    return [
        ["External developer review", "Strong if focused", "The system has real software depth when it combines verified college information, admission engagement, inquiry tracking, staff escalation, and auditable content management."],
        ["Professor review", "Suitable final-year scope", "The project is strong for a Computer Engineering final-year submission if the MVP is built completely and measurable claims are not exaggerated."],
        ["Main risk", "Scope expansion", "Trying to fully automate every university service can weaken the work. Keep the MVP focused on admission engagement and verified information support."],
        ["Recommended positioning", "University-ready SaaS style platform", "Build for one institution in the prototype while designing the schema and roles so multiple colleges under a university can use it later."],
    ]


def prd():
    doc = setup_doc(
        "CampusGuard Product Requirements Document",
        "AI powered university admission engagement and verified information support platform",
    )
    add_meta_table(doc, [
        ["Project name", PROJECT_FULL_NAME],
        ["Audience", "Final-year project team, guide, external reviewer, and implementation team"],
        ["Document purpose", "Define the product problem, users, MVP scope, outcomes, and acceptance criteria."],
        ["Recommended MVP stance", "One college prototype, university-level multi-institution design."],
    ])

    doc.add_heading("1 Executive Summary", level=1)
    doc.add_paragraph(
        "CampusGuard is a university-level admission engagement and verified information support platform. "
        "It helps prospective students and parents understand admissions, courses, fees, required documents, "
        "scholarships, hostel, transport, placements, campus facilities, and college environment without repeated physical visits. "
        "It also helps the institution manage inquiries, present verified information, route difficult questions to staff, and improve admission follow-up."
    )
    doc.add_paragraph(
        "The strongest version is not a generic chatbot. The chatbot is the user-facing interface, while the product value comes from approved information, role-based access, inquiry tracking, staff workflows, evidence display, and auditability."
    )

    doc.add_heading("2 Senior Review Verdict", level=1)
    add_table(doc, ["Review Area", "Verdict", "Comment"], common_review_points())

    doc.add_heading("3 Problem Statement", level=1)
    add_bullets(doc, [
        "Prospective students and parents often need repeated calls or visits to understand admission procedures and college facilities.",
        "Out-of-town families spend time and money visiting campus for basic information before deciding whether to apply.",
        "Admission staff answer repeated questions manually, which reduces time available for serious counseling and document verification.",
        "College information is spread across notices, brochures, websites, staff messages, PDFs, and informal communication.",
        "A normal chatbot can give wrong or outdated answers unless its information is governed and verified.",
    ])

    doc.add_heading("4 Target Users", level=1)
    add_table(doc, ["User", "Primary Needs", "MVP Access"], [
        ["Prospective student", "Courses, eligibility, documents, fees, scholarships, facilities, application guidance", "Ask chatbot, view verified answers, save checklist, submit inquiry"],
        ["Parent", "Fees, safety, hostel, transport, college environment, placements, staff contact, visit planning", "Ask chatbot, view public information, submit inquiry"],
        ["Admission staff or faculty", "Resolve escalated questions, update admission information, follow up with leads", "Staff dashboard, inquiry queue, content suggestions"],
        ["College admin", "Manage users, content, departments, inquiries, reports, and audit logs", "Admin dashboard and settings"],
        ["University admin", "Manage multiple colleges and compare inquiry/admission engagement across institutions", "University dashboard in planned scale-up"],
    ])

    doc.add_heading("5 Goals", level=1)
    add_numbered(doc, [
        "Reduce unnecessary physical visits for admission-related information.",
        "Improve trust by answering from approved college and university content.",
        "Help applicants and parents make informed admission decisions.",
        "Help the college manage inquiries and follow-up professionally.",
        "Design the system so it can scale from one college to multiple colleges under a university.",
    ])

    doc.add_heading("6 MVP Scope", level=1)
    add_table(doc, ["MVP Module", "Included", "Why It Matters"], [
        ["AI information assistant", "Admission, courses, fees, documents, scholarships, facilities, placements, hostel, transport, campus life", "Main user value"],
        ["Verified content repository", "Admin-managed FAQs, notices, policies, documents, and source links", "Prevents unsupported answers"],
        ["Lead and inquiry capture", "Interested course, student details, parent details, city, contact preference, question history", "Helps institution follow up"],
        ["Staff escalation", "Route unanswered or sensitive questions to admission staff", "Keeps human support in the process"],
        ["Document checklist guidance", "Personalized checklist based on course and category", "Reduces admission confusion"],
        ["Role-based dashboards", "Applicant/parent view, staff queue, admin reporting", "Makes the product usable by all roles"],
        ["Audit and analytics", "Question logs, content usage, unresolved topics, staff response status", "Supports improvement and viva explanation"],
    ])

    doc.add_heading("7 User Stories", level=1)
    add_table(doc, ["ID", "As a", "I want", "So that"], [
        ["US-01", "prospective student", "to ask about eligibility and courses", "I can decide whether to apply"],
        ["US-02", "parent", "to understand fees, hostel, transport, and safety", "I can evaluate the college without an unnecessary visit"],
        ["US-03", "applicant", "to get a document checklist", "I can prepare before visiting or submitting documents"],
        ["US-04", "staff member", "to receive escalated inquiries", "I can answer complex cases quickly"],
        ["US-05", "admin", "to approve and update official content", "the chatbot does not answer from outdated data"],
        ["US-06", "university admin", "to onboard colleges later", "the same platform can serve multiple institutions"],
    ])

    doc.add_heading("8 Success Metrics", level=1)
    add_table(doc, ["Metric", "Target for MVP", "Measurement"], [
        ["Answer usefulness", "80 percent positive feedback on tested queries", "User feedback after answers"],
        ["Verified answer rate", "90 percent of official answers linked to approved content", "Answer logs"],
        ["Escalation quality", "95 percent of unsupported/sensitive queries routed or clarified", "Staff review"],
        ["Inquiry capture", "All interested users can submit an inquiry with contact and course interest", "Functional test"],
        ["Response time", "Most FAQ and retrieval answers under 5 seconds in demo setup", "Backend timing logs"],
        ["Admin manageability", "Admin can add, update, approve, and deactivate content without code change", "Acceptance test"],
    ])

    doc.add_heading("9 Assumptions", level=1)
    add_bullets(doc, [
        "The prototype will use one college dataset but include college_id or institution_id for future university scaling.",
        "Real admission data and student documents will be synthetic unless the institution approves real data use.",
        "Final admission approval, fee payment, and document verification remain with authorized staff.",
        "The first demo will use text-based chat; voice and multilingual support can be future features.",
    ])

    doc.add_heading("10 Risks and Mitigations", level=1)
    add_table(doc, ["Risk", "Impact", "Mitigation"], [
        ["Over-broad scope", "Incomplete implementation", "Keep MVP focused on admission engagement and verified information"],
        ["Wrong AI answer", "Loss of trust", "Answer only from approved content and show evidence where available"],
        ["No real college data", "Weak demo", "Use realistic synthetic data and admin upload workflow"],
        ["Privacy mistake", "Security issue", "Use RBAC, avoid real documents in prototype, log sensitive actions"],
        ["Poor staff adoption", "Low usefulness", "Build staff dashboard and inquiry workflow, not only chatbot UI"],
    ])

    doc.add_heading("11 Out of Scope", level=1)
    add_bullets(doc, [
        "Final admission approval by AI.",
        "Online fee payment execution.",
        "Full university ERP replacement.",
        "Real student document verification without institutional approval.",
        "Facial recognition, biometric verification, or surveillance features.",
        "Fully automated scholarship awarding.",
    ])

    doc.add_heading("12 Acceptance Criteria", level=1)
    add_numbered(doc, [
        "A prospective student or parent can ask admission and campus information questions.",
        "The system answers from approved content or asks for clarification when information is missing.",
        "The system captures an inquiry and routes it to staff.",
        "Staff can view, filter, assign, and respond to inquiries.",
        "Admin can manage college information, FAQs, documents, departments, courses, fees, and facilities.",
        "The database supports future multi-college deployment through institution-level separation.",
        "Security rules prevent public users from accessing admin-only information.",
    ])
    return save(doc, "CampusGuard_PRD.docx")


def srs():
    doc = setup_doc(
        "CampusGuard Software Requirements Specification",
        "Detailed functional and non functional requirements for implementation",
    )
    add_meta_table(doc, [
        ["Project name", PROJECT_FULL_NAME],
        ["System type", "Web application with AI-assisted information support and staff dashboards"],
        ["Primary users", "Prospective students, parents, admission staff, college admins, university admins"],
        ["Primary goal", "Reduce admission uncertainty and unnecessary visits while improving inquiry handling."],
    ])

    doc.add_heading("1 Scope", level=1)
    doc.add_paragraph(
        "This SRS defines testable requirements for a focused MVP of CampusGuard. The system will provide verified admission and college information, manage inquiries, support staff follow-up, and keep the architecture ready for multiple institutions."
    )

    doc.add_heading("2 User Roles and Permissions", level=1)
    add_table(doc, ["Role", "Permissions", "Restrictions"], [
        ["Guest", "Ask public questions, browse college information, start inquiry", "Cannot view inquiry status unless verified by link or login"],
        ["Prospective student", "Manage own profile, ask questions, save checklist, submit and track inquiry", "Cannot access admin or other applicant records"],
        ["Parent", "Ask public questions, submit inquiry, optionally link to applicant profile", "Cannot view private applicant data without consent"],
        ["Admission staff", "View assigned inquiries, answer escalations, propose content updates", "Cannot manage system roles unless also admin"],
        ["College admin", "Manage college content, courses, fees, staff, reports, and approved answers", "Can access only assigned college"],
        ["University admin", "Manage institutions, global settings, and cross-college reports", "Should not read private inquiry content unless policy permits"],
    ])

    doc.add_heading("3 Functional Requirements", level=1)
    add_table(doc, ["ID", "Requirement", "Acceptance Test"], [
        ["FR-01", "The system shall allow public users to ask admission and college information questions.", "A guest submits a question and receives an answer, clarification, or escalation option."],
        ["FR-02", "The system shall classify queries into domains such as courses, fees, documents, scholarships, hostel, transport, placements, facilities, and campus life.", "Each sample query is stored with a domain label."],
        ["FR-03", "The system shall answer official questions only from approved content records.", "A draft or disabled content record is not used in official responses."],
        ["FR-04", "The system shall show source title or content reference for important factual answers.", "Answers about fees, dates, documents, and eligibility display source information."],
        ["FR-05", "The system shall ask a clarification question when course, category, year, or campus is required.", "Missing context returns a clarification rather than a guessed answer."],
        ["FR-06", "The system shall capture inquiries with name, phone/email, location, course interest, role, and message.", "A valid inquiry appears in the staff dashboard."],
        ["FR-07", "The system shall route inquiries by college, department, course, or domain.", "A scholarship query goes to the configured admission or scholarship contact."],
        ["FR-08", "The system shall let staff respond to assigned inquiries and update status.", "Staff changes status from new to in progress to resolved."],
        ["FR-09", "The system shall allow admins to manage courses, departments, fees, facilities, documents, FAQs, and notices.", "Admin creates and approves a content record visible to the assistant."],
        ["FR-10", "The system shall keep an audit record for content changes, staff responses, and role changes.", "Audit table records actor, action, target, timestamp, and result."],
        ["FR-11", "The system shall support institution_id or college_id on institution-owned records.", "Two colleges can store separate course and fee records."],
        ["FR-12", "The system shall provide reports for inquiry volume, top questions, unresolved questions, and conversion status.", "Admin dashboard displays metrics from sample data."],
    ])

    doc.add_heading("4 Business Rules", level=1)
    add_bullets(doc, [
        "Official admission dates, fees, documents, eligibility, scholarships, and seat information must come from approved records.",
        "The assistant must not promise admission, scholarship approval, fee waiver, hostel allocation, or placement outcome.",
        "Sensitive inquiries must be routed to staff rather than answered automatically.",
        "Each college can manage its own content; university admin can configure shared templates and global categories.",
        "Expired notices must not be shown as current information unless the user asks for historical context.",
    ])

    doc.add_heading("5 Data Requirements", level=1)
    add_table(doc, ["Entity", "Key Fields", "Purpose"], [
        ["Institution", "id, name, type, city, status", "University or college tenant"],
        ["User", "id, name, email, phone, role, institution_id, status", "Identity and access control"],
        ["Course", "id, institution_id, name, department, duration, eligibility, intake", "Course discovery"],
        ["FeeStructure", "id, institution_id, course_id, academic_year, components, due_dates, status", "Fee guidance"],
        ["ContentItem", "id, institution_id, title, category, body, source_type, status, approved_by", "Approved answer source"],
        ["Document", "id, institution_id, title, file_path, extracted_text, status, checksum", "Uploaded college documents"],
        ["Inquiry", "id, institution_id, name, contact, role, city, course_interest, message, status, owner_id", "Lead and support tracking"],
        ["ChatMessage", "id, session_id, user_id, message, response, sources, outcome", "Conversation history"],
        ["AuditLog", "id, actor_id, action, entity_type, entity_id, timestamp, metadata", "Accountability"],
    ])

    doc.add_heading("6 Validations", level=1)
    add_bullets(doc, [
        "Email must use valid email format when provided.",
        "Phone number must meet configured country or institution format.",
        "Course interest must match an active course when selected from the course list.",
        "Required inquiry fields must be present before creating a staff task.",
        "Fee amounts and deadlines must be admin-entered, non-negative, and tied to an academic year.",
        "Uploaded files must be limited by type, size, and malware-scan status before processing.",
    ])

    doc.add_heading("7 Authentication and Authorization", level=1)
    add_bullets(doc, [
        "Use JWT-based sessions for the student, parent, staff, admin, and university admin roles.",
        "Passwords must be hashed using a secure one-way password hashing algorithm.",
        "All admin and staff APIs must verify the role server-side.",
        "College admin and staff access must be scoped to their institution_id.",
        "Public chat must not expose admin-only notes, private inquiries, or unpublished content.",
    ])

    doc.add_heading("8 Error Handling and Edge Cases", level=1)
    add_table(doc, ["Case", "Expected Behavior"], [
        ["No matching approved answer", "Ask for details or create inquiry; do not fabricate."],
        ["Conflicting fee/date information", "Show uncertainty and route to staff."],
        ["User asks for final admission guarantee", "Explain that final decision belongs to authorized admission staff."],
        ["Database unavailable", "Show friendly failure and log technical error."],
        ["Staff response delayed", "Keep inquiry status visible and allow reminder/escalation."],
        ["Prompt injection inside uploaded content", "Treat document text as data and do not follow embedded instructions."],
    ])

    doc.add_heading("9 Non Functional Requirements", level=1)
    add_table(doc, ["Category", "Requirement"], [
        ["Security", "Protect staff/admin APIs, validate input, store secrets in environment variables, and avoid real personal data in demo."],
        ["Performance", "Common chat answers should return within 5 seconds in the demo environment."],
        ["Reliability", "Important facts must be answerable from approved content or escalated."],
        ["Usability", "The interface must be responsive, clear, and usable by parents and students with limited technical comfort."],
        ["Maintainability", "Separate frontend, backend, database migrations, services, and documentation."],
        ["Auditability", "Every content approval, inquiry status change, and staff response must be traceable."],
    ])

    doc.add_heading("10 Acceptance Criteria", level=1)
    add_numbered(doc, [
        "All roles can log in or use allowed public functions.",
        "Official answers are restricted to approved content.",
        "Inquiry capture and staff resolution workflow works end to end.",
        "Admin can manage the main information categories without code changes.",
        "Multi-institution fields exist and are enforced for admin/staff data.",
        "Security tests confirm unauthorized users cannot access staff/admin APIs.",
    ])
    return save(doc, "CampusGuard_SRS.docx")


def architecture():
    doc = setup_doc(
        "CampusGuard System Architecture Document",
        "Practical architecture for scalable university admission engagement and verified information support",
    )
    add_meta_table(doc, [
        ["Project name", PROJECT_FULL_NAME],
        ["Architecture style", "Modular web application with REST APIs and AI-assisted retrieval"],
        ["Deployment target", "Local/demo deployment first, cloud-ready later"],
        ["Scale target", "Single college prototype with university-level multi-college readiness"],
    ])

    doc.add_heading("1 Architecture Decision", level=1)
    doc.add_paragraph(
        "CampusGuard should use a practical three-layer architecture: React frontend, FastAPI backend, and relational database. "
        "The AI assistant should sit behind backend services, not directly access the database. This keeps the implementation understandable for students while still supporting security, governance, and future scale."
    )

    doc.add_heading("2 Recommended Tech Stack", level=1)
    add_table(doc, ["Layer", "Recommended Choice", "Reason"], [
        ["Frontend", "React with Tailwind CSS", "Fast development, clean dashboards, responsive UI"],
        ["Backend", "Python FastAPI", "Clear REST APIs, typed request models, good AI integration"],
        ["Database", "PostgreSQL for final design or MySQL if college stack requires it", "Relational data, reporting, indexing, multi-institution schema"],
        ["Search", "PostgreSQL full-text search for MVP; vector search optional", "Good enough for verified FAQ and document retrieval"],
        ["AI", "Backend-managed LLM or API-based model", "Natural-language responses with controlled prompts"],
        ["Auth", "JWT with hashed passwords", "Simple, explainable, standard for student projects"],
        ["Files", "Local storage for demo; object storage later", "Simple prototype with future cloud path"],
        ["Deployment", "Docker-ready services", "Portable for viva and cloud deployment"],
    ])

    doc.add_heading("3 Logical Components", level=1)
    add_table(doc, ["Component", "Responsibility"], [
        ["Web Client", "Applicant, parent, staff, and admin screens"],
        ["API Gateway or Backend Router", "Routes requests and applies authentication middleware"],
        ["Auth and RBAC Service", "Login, sessions, roles, tenant scoping"],
        ["Institution Service", "University, college, department, course, and facility configuration"],
        ["Content Service", "Approved FAQs, notices, policies, and document records"],
        ["Retrieval Service", "Search approved content and return candidate evidence"],
        ["AI Answer Service", "Generate user-friendly answer from retrieved approved content"],
        ["Inquiry Service", "Create, assign, respond to, and track inquiries"],
        ["Admin Reporting Service", "Dashboards for engagement, questions, unresolved cases, and staff workload"],
        ["Audit Service", "Record content, role, inquiry, and sensitive actions"],
    ])

    doc.add_heading("4 Data Flow", level=1)
    add_numbered(doc, [
        "User selects role or continues as guest.",
        "User asks a question or browses an admission category.",
        "Backend classifies the query domain and tenant context.",
        "Retrieval service searches approved content for the selected college or university.",
        "AI answer service drafts a response only from returned content.",
        "Verifier checks whether the response has enough source support for important facts.",
        "System returns answer, asks clarification, or creates an inquiry for staff.",
        "Inquiry and answer metadata are stored for reporting and audit.",
    ])

    doc.add_heading("5 Database Architecture", level=1)
    add_table(doc, ["Table Group", "Tables"], [
        ["Identity", "users, roles, user_roles, sessions"],
        ["Institutions", "institutions, colleges, departments, campuses"],
        ["Admissions", "courses, admission_cycles, fee_structures, scholarships, document_checklists"],
        ["Content", "content_items, documents, document_chunks, content_approvals"],
        ["Chat", "chat_sessions, chat_messages, answer_sources"],
        ["Inquiries", "inquiries, inquiry_assignments, inquiry_messages, inquiry_status_history"],
        ["Reporting", "question_analytics, lead_sources, conversion_stages"],
        ["Audit", "audit_logs, security_events"],
    ])

    doc.add_heading("6 API Structure", level=1)
    add_table(doc, ["API Group", "Example Endpoints", "Allowed Roles"], [
        ["Auth", "POST /auth/login, POST /auth/register, POST /auth/logout", "Public, all logged-in users"],
        ["Institutions", "GET /institutions, GET /institutions/{id}/profile", "Public, admin for writes"],
        ["Courses", "GET /courses, POST /admin/courses", "Public read, admin write"],
        ["Content", "GET /content/search, POST /admin/content, PUT /admin/content/{id}/approve", "Public read approved, admin write"],
        ["Chat", "POST /chat/query, GET /chat/sessions", "Public limited, logged-in history"],
        ["Inquiries", "POST /inquiries, GET /staff/inquiries, PATCH /staff/inquiries/{id}", "Public create, staff/admin manage"],
        ["Reports", "GET /admin/reports/overview", "Admin, university admin"],
        ["Audit", "GET /admin/audit-logs", "Admin only"],
    ])

    doc.add_heading("7 Security Design", level=1)
    add_bullets(doc, [
        "All write APIs require authentication and server-side authorization.",
        "Admin and staff data access must be scoped by institution_id.",
        "The AI service must not use unpublished content for official answers.",
        "Prompt inputs, uploaded content, and OCR text must be treated as data, not instructions.",
        "Secrets such as JWT keys and model API keys must be stored in environment variables.",
        "Files must be validated by type and size; production should add malware scanning.",
    ])

    doc.add_heading("8 Deployment and Monitoring", level=1)
    add_bullets(doc, [
        "Use separate frontend and backend services for development.",
        "Use database migrations for schema changes.",
        "Containerize frontend, backend, and database for reproducible demo deployment.",
        "Log request ID, user role, institution_id, API route, outcome, and error code.",
        "Track common unanswered questions so admins can add missing content.",
    ])

    doc.add_heading("9 Scalability Plan", level=1)
    add_table(doc, ["Stage", "Architecture Support"], [
        ["Single college demo", "One institution record with sample courses, fees, and content"],
        ["Multi-college university", "institution_id on content, users, inquiries, courses, fees, reports"],
        ["Higher traffic", "Caching for approved content, pagination, background document processing"],
        ["Advanced AI retrieval", "Add embeddings and vector index after basic verified content search works"],
    ])

    doc.add_heading("10 Architecture Acceptance Criteria", level=1)
    add_numbered(doc, [
        "Frontend communicates only through documented backend APIs.",
        "AI service cannot directly read or write database tables.",
        "All institution-owned data includes tenant scoping.",
        "Approved content controls official answers.",
        "Staff inquiry flow and admin content flow are auditable.",
    ])
    return save(doc, "CampusGuard_Architecture.docx")


def uiux():
    doc = setup_doc(
        "CampusGuard UI UX Document",
        "Implementation ready user experience for applicants parents staff and admins",
    )
    add_meta_table(doc, [
        ["Project name", PROJECT_FULL_NAME],
        ["Design goal", "Simple, trustworthy, responsive, and easy for parents and students to use"],
        ["Primary screens", "Public home, chatbot, admission explorer, inquiry flow, staff dashboard, admin dashboard"],
        ["Tone", "Professional college support, not marketing-only and not over-technical"],
    ])

    doc.add_heading("1 Design Principles", level=1)
    add_bullets(doc, [
        "Make the first screen useful immediately: search, ask, and explore admission information.",
        "Use verified information labels for official facts such as dates, fees, documents, and eligibility.",
        "Keep parent and applicant language simple and action-oriented.",
        "Show clear next steps after every answer.",
        "Avoid overcrowded dashboards; prioritize common admission tasks.",
        "Design for mobile first because many users will access the system from phones.",
    ])

    doc.add_heading("2 User Journeys", level=1)
    add_table(doc, ["Journey", "Steps", "Outcome"], [
        ["Prospective student admission query", "Open site, choose college, ask course/eligibility question, receive verified answer, save checklist, submit inquiry", "Student understands next step"],
        ["Parent college evaluation", "Browse facilities, fees, hostel, transport, safety, placements, ask question, request callback", "Parent can decide whether a visit is needed"],
        ["Staff inquiry resolution", "Open staff dashboard, view assigned inquiries, read chat context, respond, mark status", "Question is resolved with staff accountability"],
        ["Admin content update", "Open admin panel, add notice/FAQ/document, approve content, review unanswered questions", "Assistant improves without code changes"],
    ])

    doc.add_heading("3 Navigation Model", level=1)
    add_table(doc, ["Area", "Navigation Items"], [
        ["Public", "Home, Ask CampusGuard, Courses, Admissions, Fees, Scholarships, Facilities, Contact"],
        ["Applicant", "Dashboard, My Questions, Document Checklist, Saved Courses, Inquiry Status"],
        ["Parent", "Dashboard, Fees, Facilities, Hostel and Transport, Questions, Contact Staff"],
        ["Staff", "Dashboard, Assigned Inquiries, Escalated Questions, Content Suggestions, Follow-ups"],
        ["Admin", "Dashboard, Courses, Fees, Content, Documents, Users, Reports, Audit Logs, Settings"],
        ["University admin", "Institutions, Colleges, Cross-college Reports, Global Categories, Admin Users"],
    ])

    doc.add_heading("4 Screen List", level=1)
    add_table(doc, ["Screen", "Primary Components", "Important States"], [
        ["Home", "College selector, ask bar, quick categories, admission CTA, contact CTA", "Loading institutions, no college selected"],
        ["Ask CampusGuard", "Chat panel, source panel, suggested questions, inquiry button", "Typing, no answer, needs clarification, escalated"],
        ["Course Explorer", "Course cards/table, filters, eligibility summary, apply interest", "No courses, filter empty, course inactive"],
        ["Admission Guide", "Timeline, steps, document checklist, fee links, staff contact", "Deadline passed, missing category"],
        ["Facilities", "Hostel, transport, labs, library, campus life, placements", "Information unavailable, request callback"],
        ["Inquiry Form", "Name, role, city, course, contact, question, consent checkbox", "Validation errors, submitted, duplicate inquiry"],
        ["Applicant Dashboard", "Saved checklist, recent questions, inquiry status, recommended next steps", "New user, no inquiry yet"],
        ["Staff Dashboard", "Queue, filters, inquiry detail, response editor, status controls", "No assigned inquiries, overdue items"],
        ["Admin Dashboard", "Metrics, unanswered questions, content status, staff performance", "No data, export loading"],
        ["Content Manager", "FAQ/editor, document upload, approval state, source details", "Draft, pending approval, rejected, published"],
    ])

    doc.add_heading("5 Core Interactions", level=1)
    add_bullets(doc, [
        "Ask bar should accept free-text questions and optional college/course context.",
        "Important answers should show source title and last updated date.",
        "If the assistant cannot answer, show a staff inquiry option immediately.",
        "Forms should use clear field labels, inline validation, and simple error messages.",
        "Staff queue should support status filters, course filters, priority, and assignment.",
        "Admin content update should separate draft, review, approved, and inactive states.",
    ])

    doc.add_heading("6 Form Requirements", level=1)
    add_table(doc, ["Form", "Fields", "Validation"], [
        ["Inquiry", "Name, role, city, phone/email, course interest, message, consent", "Contact required, message required, consent required"],
        ["Course", "Name, department, duration, eligibility, intake, fees link, status", "Name and department required; intake non-negative"],
        ["Fee", "Course, academic year, components, due date, notes, status", "Amount non-negative; due date valid"],
        ["Content item", "Title, category, body, source, institution, status", "Title, category, and body required"],
        ["Staff response", "Message, status, next follow-up date", "Message required when resolving inquiry"],
    ])

    doc.add_heading("7 Loading Error and Empty States", level=1)
    add_table(doc, ["State", "UI Behavior"], [
        ["Loading chat answer", "Show compact loader and keep previous messages visible."],
        ["No verified answer", "Say information is not available and offer inquiry creation."],
        ["Network error", "Show retry button and preserve typed question."],
        ["Empty inquiry queue", "Show clean empty state with filters and no false error."],
        ["No course results", "Show clear message and reset-filter option."],
        ["Unauthorized", "Show access denied and return to safe dashboard."],
    ])

    doc.add_heading("8 Responsive Behavior", level=1)
    add_bullets(doc, [
        "Mobile: single-column layout, sticky ask button, collapsible filters, large touch targets.",
        "Tablet: two-column content where useful, chat and source panel stacked if space is tight.",
        "Desktop: dashboard side navigation, chat plus source panel, staff inquiry detail split view.",
        "All text must wrap safely and buttons must not resize awkwardly with long labels.",
    ])

    doc.add_heading("9 Accessibility", level=1)
    add_bullets(doc, [
        "Use semantic headings and form labels.",
        "Maintain sufficient contrast for text, buttons, and statuses.",
        "Support keyboard navigation for forms, chat input, filters, and modals.",
        "Do not rely only on color for statuses; include text labels.",
        "Use readable font sizes and line spacing for parents and first-time users.",
    ])

    doc.add_heading("10 Visual System", level=1)
    add_table(doc, ["Design Token", "Recommendation"], [
        ["Typography", "Aptos, Inter, or system sans; 16px base in web UI"],
        ["Primary color", "Deep institutional blue for primary actions"],
        ["Accent color", "Fresh green or teal for success and engagement highlights"],
        ["Warning color", "Amber for pending or clarification states"],
        ["Error color", "Red for validation and access errors"],
        ["Spacing", "8px spacing system with clear section separation"],
        ["Cards", "Use only for repeated items such as courses, inquiries, and dashboard metrics"],
    ])

    doc.add_heading("11 UI Acceptance Criteria", level=1)
    add_numbered(doc, [
        "A first-time parent can find fees, hostel, transport, and contact support within three actions.",
        "A student can ask a question and submit an inquiry from the same flow.",
        "Staff can resolve an inquiry without opening unrelated admin screens.",
        "Admin can publish content and see whether it is available to the assistant.",
        "Mobile layouts have no overlapping text, clipped buttons, or unreadable tables.",
    ])
    return save(doc, "CampusGuard_UI_UX.docx")


def development_plan():
    doc = setup_doc(
        "CampusGuard Development Plan",
        "Roadmap milestones priorities dependencies and definition of done",
    )
    add_meta_table(doc, [
        ["Project name", PROJECT_FULL_NAME],
        ["Planning stance", "Build one complete vertical slice first, then expand categories"],
        ["Recommended team approach", "Module by module with working demo after each milestone"],
        ["MVP deadline goal", "A complete explainable prototype before adding advanced AI features"],
    ])

    doc.add_heading("1 Build Strategy", level=1)
    doc.add_paragraph(
        "The team should build CampusGuard as a working admission engagement platform before adding advanced research features. "
        "The first complete vertical slice should be: approved content -> chatbot answer -> inquiry escalation -> staff response -> admin report. "
        "Once that is stable, add document upload, better retrieval, analytics, and multi-college configuration."
    )

    doc.add_heading("2 Priority Levels", level=1)
    add_table(doc, ["Priority", "Meaning", "Examples"], [
        ["P0", "Required for MVP demo", "Auth, roles, approved content, chat, inquiry workflow, admin dashboard"],
        ["P1", "Important polish and completeness", "Document upload, source display, reports, filters, email-style notifications"],
        ["P2", "Advanced or future", "Vector search, OCR, multilingual, voice, CRM integrations, predictive conversion"],
    ])

    doc.add_heading("3 Milestone Roadmap", level=1)
    add_table(doc, ["Milestone", "Work", "Exit Criteria"], [
        ["M1 Project setup", "Create frontend, backend, database, environment files, README", "App runs locally and health API works"],
        ["M2 Database foundation", "Create tables for users, roles, institutions, courses, content, inquiries, audit", "Migrations run and seed data loads"],
        ["M3 Authentication and RBAC", "Login, JWT, role middleware, route protection", "Each role lands on correct dashboard"],
        ["M4 Admin content management", "CRUD for courses, fees, facilities, FAQs, notices, approval status", "Approved content appears in public views"],
        ["M5 Chat and retrieval MVP", "Question API, domain detection, approved content search, answer formatting", "Common questions return verified answers or clarification"],
        ["M6 Inquiry workflow", "Inquiry form, assignment, staff response, status history", "User can submit and staff can resolve inquiry"],
        ["M7 Dashboards and reports", "Admin metrics, staff queue, top questions, unanswered topics", "Reports reflect seed/demo activity"],
        ["M8 Document support", "Upload PDFs/DOCX, store metadata, extract text where available", "Admin can attach documents to content"],
        ["M9 QA and security", "Validation, error handling, unauthorized access tests, edge cases", "Critical flows pass test checklist"],
        ["M10 Deployment and viva prep", "Docker/cloud setup, demo data, screenshots, final docs, presentation support", "Demo can be run repeatably"],
    ])

    doc.add_heading("4 Dependencies", level=1)
    add_table(doc, ["Task", "Depends On"], [
        ["Chat answer API", "Approved content schema and seed data"],
        ["Staff dashboard", "Inquiry table and authentication"],
        ["Admin reports", "Chat logs and inquiry status history"],
        ["Multi-college support", "Institution table and tenant scoping from the start"],
        ["Document upload", "Admin auth and content management"],
        ["Advanced AI retrieval", "Stable approved content repository"],
        ["Deployment", "Environment configuration and database migrations"],
    ])

    doc.add_heading("5 Suggested Folder Structure", level=1)
    doc.add_paragraph("Use a structure that separates frontend, backend, database, and documentation clearly.")
    add_table(doc, ["Folder", "Purpose"], [
        ["frontend/", "React application, pages, components, API client, styles"],
        ["backend/", "FastAPI application, routes, services, schemas, auth, tests"],
        ["database/", "SQL migrations, seed data, ER diagram"],
        ["docs/", "PRD, SRS, architecture, UI/UX, development plan, viva notes"],
        ["uploads/", "Local development file storage; ignored from git for private files"],
        ["README.md", "Setup, run commands, project overview"],
    ])

    doc.add_heading("6 Testing Plan", level=1)
    add_table(doc, ["Test Area", "What to Test"], [
        ["Unit tests", "Validation helpers, role checks, domain classification, content filters"],
        ["API tests", "Auth, content CRUD, chat, inquiries, reports"],
        ["Security tests", "Unauthorized access, tenant isolation, unpublished content leakage"],
        ["UI tests", "Login, chat flow, inquiry form, staff queue, admin content"],
        ["Data tests", "Seed data, migrations, required fields, foreign keys"],
        ["Demo tests", "End-to-end flow from parent question to staff resolution"],
    ])

    doc.add_heading("7 Definition of Done", level=1)
    add_bullets(doc, [
        "Feature has frontend UI, backend API, validation, and database persistence where needed.",
        "Role permissions are enforced server-side.",
        "Empty, loading, error, and success states are handled.",
        "Important user actions are logged in audit or status history.",
        "At least one positive and one negative test case is documented.",
        "The feature is explainable in viva with data flow and database tables.",
    ])

    doc.add_heading("8 MVP Build Order", level=1)
    add_numbered(doc, [
        "Start with database design and seed data for one college.",
        "Build login and role dashboards.",
        "Build admin content management for courses, fees, documents, FAQs, facilities, and scholarships.",
        "Build public pages and chatbot retrieval from approved content.",
        "Build inquiry capture and staff response workflow.",
        "Add reports and top unanswered questions.",
        "Add document upload and extraction support.",
        "Add security tests, UI polish, and deployment.",
    ])

    doc.add_heading("9 Features to Avoid Until Core MVP Works", level=1)
    add_bullets(doc, [
        "Facial recognition or biometric systems.",
        "Payment gateway integration.",
        "Real admission decision automation.",
        "Full ERP or SIS integration.",
        "Predictive admission conversion scoring.",
        "Voice and multilingual mode before the text workflow is stable.",
        "Too many separate service modules that become shallow CRUD pages.",
    ])

    doc.add_heading("10 Final External Review Note", level=1)
    doc.add_paragraph(
        "CampusGuard is a strong final-year project if the team demonstrates a complete, focused product: verified admission information, applicant and parent engagement, staff follow-up, admin-controlled content, multi-college readiness, and clear security boundaries. "
        "It becomes weak only if it is presented as a generic chatbot or if the team adds many unfinished modules without a working core."
    )
    return save(doc, "CampusGuard_Development_Plan.docx")


def main():
    paths = [prd(), srs(), architecture(), uiux(), development_plan()]
    for path in paths:
        print(path.resolve())


if __name__ == "__main__":
    main()
